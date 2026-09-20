"""Editable PPTX plus PDF/contact-sheet previews; all charts are native shapes."""
import json, math
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'presentations'
E=json.loads((OUT/'evidence.json').read_text())
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
INK='20282D';MUTED='657078';ACC='365C6D';LIGHT='F1F3F4';LINE='DCE1E4';GRAY='B9C1C6';WHITE='FFFFFF'
SCALE=120; pages=[]; overflows=[]; current=None; canvas=None;draw=None
FONT='C:/Windows/Fonts/msjh.ttc';BOLD='C:/Windows/Fonts/msjhbd.ttc'
def rgb(c):return RGBColor.from_string(c)
def rect(x,y,w,h,fill=LIGHT,line=None):
    s=current.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));s.fill.solid();s.fill.fore_color.rgb=rgb(fill)
    if line:s.line.color.rgb=rgb(line)
    else:s.line.fill.background()
    draw.rectangle([x*SCALE,y*SCALE,(x+w)*SCALE,(y+h)*SCALE],fill='#'+fill,outline='#'+line if line else None)
def text(x,y,w,h,t,size=18,color=INK,bold=False):
    font=ImageFont.truetype(BOLD if bold else FONT,round(size*SCALE/72))
    lines=[]
    for paragraph in t.split('\n'):
        line=''
        for ch in paragraph:
            if line and draw.textlength(line+ch,font=font)>w*SCALE-3:lines.append(line);line=''
            line+=ch
        lines.append(line)
    step=size/72*1.33
    if len(lines)*step>h+.06:overflows.append((len(prs.slides),t[:32],len(lines)*step,h))
    tb=current.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Microsoft JhengHei';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
        p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=1.2
        for run in p.runs:
            ea=OxmlElement('a:ea');ea.set('typeface','Microsoft JhengHei');run._r.get_or_add_rPr().append(ea)
        draw.text((x*SCALE,(y+i*step)*SCALE),line,font=font,fill='#'+color)
def slide(title,kicker='DESIGN RATIONALE',source='',notes=''):
    global current,canvas,draw
    if canvas is not None:pages.append(canvas)
    current=prs.slides.add_slide(prs.slide_layouts[6]);canvas=Image.new('RGB',(1600,900),'white');draw=ImageDraw.Draw(canvas)
    text(.55,.3,11,.25,kicker,10,ACC,True);text(.55,.78,12.2,.7,title,29,bold=True)
    rect(.55,6.99,12.2,.012,LINE);text(.55,7.12,11.6,.25,source,9,MUTED);text(12.35,7.1,.45,.3,str(len(prs.slides)).zfill(2),11,MUTED)
    current.notes_slide.notes_text_frame.text=notes+'\n來源：'+source
def card(x,y,w,title,body):
    rect(x,y,w,2.0);text(x+.2,y+.2,w-.4,.45,title,21,ACC,True);text(x+.2,y+.85,w-.4,1.0,body,17)
def arrow(x,y):text(x,y,.5,.5,'→',25,ACC,True)
def takeaway(t):rect(.55,6.12,12.2,.57,LIGHT);text(.75,6.22,11.8,.4,t,17,ACC,True)

slide('把判斷帶進測試流程','ADVANTEST ACS RTDI  /  場景 1 + 場景 2',
      '依據 Question_20260919.pdf 與本專案實驗；2026.09.20',
      '開場：當一片 wafer 還在測試，工程師能否已經知道哪裡值得關注？本方案有兩個独立任務，分別服務工程師與測試程式。不是把兩個模型合成一個分數。')
text(.6,1.85,11.5,.75,'從「測完再分析」走向「測試中取得可用結果」',27)
card(.6,3.05,5.8,'01  發現異常','為工程師整理 wafer 狀態、失敗位置與證據')
card(6.65,3.05,6.05,'02  預測溫度','在指定測項前，向測試程式回傳預測值')
takeaway('相同資料基礎，兩個獨立任務；重點是結果能否在正確時機被使用。')

slide('問題不只在模型，而在判斷來得太晚','01  問題與目標','題目：Question_20260919.pdf，第 1–2、7 頁',
      '題目描述的是測試完成後的資料處理負擔。評分中 Gemini 運行與正確時機各 25%，合計 50%。這是設計優先順序的依據，不是我們已獲得的成績。')
for x,t,b in [(0.6,'現在的負擔','測試完成\n再整理資料、找異常'),(4.8,'希望的改變','測試進行中\n產生分類／預測'),(9,'可用的結果','人員能追查\n機台程式能接收')]:card(x,1.9,3.7,t,b)
arrow(4.38,2.65);arrow(8.58,2.65)
text(.65,4.5,3.4,.6,'50% 評分',32,ACC,True)
text(4.05,4.52,8.3,.8,'Gemini 運行 25% ＋ 正確偵測／預測時機 25%',22)
takeaway('先解決資料、時序與回傳，再討論模型複雜度。')

slide('資料量看似很大，獨立的標籤其實很少','02  資料限制','題目第 3 頁；data/TrainDataInfo.txt；artifacts/wafer_classifier/info.json',
      '排除 W2 後是 24 片，每片 80 die，共 1920 die；每 die 3036 測項。但場景一的類別是 wafer 層級，不能把每個 die 當成獨立的異常類別標籤，也不能宣稱知道哪些 die 是根因。Site 特徵目前啟用於單一模型，沒有每 site 各訓練模型。')
for x,b,s in [(0.65,'24','可用 wafers'),(4.8,'1,920','die 量測列'),(9,'3,036','每 die 測項')]:
    text(x,1.9,3.5,.8,b,44,ACC,True);text(x,2.85,3.5,.4,s,19)
rect(.6,3.7,12.1,1.75,LIGHT);text(.85,3.95,11.55,.55,'場景 1 的監督訊號只有 24 個 wafer 標籤',25,bold=True)
text(.85,4.7,11.4,.55,'17 Normal  ／  2 Low yield  ／  其餘 5 類各 1 片',22)
takeaway('按 wafer 分組驗證；不能隨機拆 die，製造虛高的成績。')

slide('場景 1：先判斷 wafer，再讓人追到 die','03  異常偵測的設計','realtime/classifier_features.py、classifier_engine.py、notifications.py',
      '現行為 L2 logistic regression，使用 44 個摘要特徵，包括 fail 比例、site差異與分布變化。TestEnd 後累積已收到且完整的資料，至少16 die開始分類。通知政策另要求至少24 die且三次一致非Normal判斷，或異常WaferEnd。這些是目前工程設定，尚非實驗證明的最佳門檻。PF Fail與wafer模型類別分開。')
for x,t,b in [(.6,'已完成的 die','實測值、PF、site\n只取目前已收到資料'),(4.8,'44 個摘要特徵','失敗比例與分布變化\nL2 Logistic Regression'),(9,'單一 wafer label','官方類別＋信心分數\n附實測位置與量測證據')]:card(x,1.85,3.7,t,b)
arrow(4.38,2.6);arrow(8.58,2.6)
text(.75,4.5,5.7,.8,'PF / Fail：測試得到的結果',23,bold=True)
text(6.8,4.5,5.6,.8,'Wafer label：模型的整體判斷',23,bold=True)
takeaway('不把 wafer 標籤貼到每顆 die；也不把統計偏離直接叫作根因。')

slide('為何先選簡單分類器？現有比較支持這個取捨','04  場景 1 的量化證據','reports/wafer_classifier/comparison.json、README.md；24 片留一片驗證',
      '固定候選：Logistic L2 19/24，Random Forest17/24，Extra Trees18/24。Normal誤報分別0/17、0/17、1/17。五個單例類別的唯一wafer留出後，訓練集中沒有該類，所以不能解讀為七類泛化準確率。模型是依可學類別平均召回與Normal誤報選擇；結果用於選型，沒有獨立最終測試集。24/24訓練內重播刻意不作效能證據。')
text(.75,1.75,5,.4,'完整 wafer 留一片驗證',19,bold=True)
for y,label,n,fp in [(2.5,'Logistic L2',19,0),(3.48,'Random Forest',17,0),(4.46,'Extra Trees',18,1)]:
    text(.75,y,2.6,.45,label,20);rect(3.5,y+.04,5.1,.28,LIGHT);rect(3.5,y+.04,5.1*n/24,.28,ACC if n==19 else GRAY)
    text(8.9,y,1.3,.4,f'{n} / 24',20,bold=True);text(10.45,y,2.1,.6,f'正常誤報 {fp}/17',17)
takeaway('支持採用較簡單的候選；五個單例異常類別的泛化能力仍未被證明。')

slide('場景 2：先問「現在拿得到什麼」，再預測','05  溫度預測的設計','Question_20260919.pdf 第 3–5 頁；scene2/train_temp_models.py；Main.flow',
      '每個sensor一個標準化Lasso，基礎前段特徵加先前sensor，排除未來測項。現行設計不使用subflow特徵。中位數補值與標準化在訓練折內估計。線上偏差修正只可使用已收到殘差；其增益尚未在此驗證表中量化。')
for x,t,b in [(.6,'前段測項','Suite / IDDQ\n作為共同輸入'),(4.8,'預測 sensor k','6 個獨立 Lasso\n可加入先前 sensor 實測'),(9,'量測 sensor k','對照預測與實測\n誤差可供後續校正')]:card(x,1.85,3.7,t,b)
arrow(4.38,2.6);arrow(8.58,2.6)
rect(.65,4.45,12,.8,LIGHT);text(.85,4.62,11.5,.45,'sensor 1 → 實測 1 → sensor 2 → 實測 2 → … → sensor 6',22,ACC,True)
takeaway('限制輸入的時間範圍，才能讓離線模型具備上線使用的可能。')

slide('溫度預測比「固定猜平均值」更有資訊','06  場景 2 的量化證據','artifacts/scene2/evaluation.json；presentations/evidence.json（同折重算基準）',
      'W2排除、24wafer、1920die；外層5折按wafer分組，同一wafer不跨訓練／測試。灰色基準為每折僅用訓練wafer的溫度平均值。Lasso結果使用已保存的相同分組評估；未套線上校正，假設前序特徵可用。這不是現場延遲、缺值或量產效能證明。MAE降低範圍63.1%至98.4%。')
text(.7,1.65,7,.4,'MAE（°C）越低越好；5 折按 wafer 分組',18)
rect(9.3,1.72,.2,.16,GRAY);text(9.65,1.63,1.5,.4,'固定平均',14)
rect(11.1,1.72,.2,.16,ACC);text(11.45,1.63,1.1,.4,'Lasso',14)
for i,r in enumerate(E['rows']):
    y=2.25+i*.54;text(.75,y,1.15,.4,'sensor '+str(r['sensor']),16)
    rect(2.2,y+.01,r['baseline_mae']/.35*6,.13,GRAY);rect(2.2,y+.19,r['model_mae']/.35*6,.13,ACC)
    text(8.6,y,2.1,.4,f"{r['baseline_mae']:.3f} → {r['model_mae']:.3f}",17)
    text(11,y,1.6,.4,f"−{r['reduction_pct']:.1f}%",18,ACC,True)
takeaway('六個 sensor 的 MAE 下降 63–98%；尚未證明現場準時回覆與線上校正增益。')

slide('場景 1 的 UI：先做判斷，再逐步看細節','07  資訊如何服務工程師','realtime/web/hierarchy.js、outcomes.js、notifications.js；建議資訊層級示意',
      '本頁是資訊架構示意，非現場截圖。第一層是目前wafer、單一label、信心與更新狀態；第二層Fail及yield；第三層空間分布與時間趨勢；點擊後才展開PID與測項。信心尚未校準，不能等同準確率。是否真的降低操作時間需做使用者任務測試。')
rect(.65,1.8,7.8,3.85,LIGHT)
text(.95,2.0,7,.5,'Wafer  →  Label  →  Confidence',24,ACC,True)
text(.95,2.7,7,.45,'Progress       Fail       Yield',21)
rect(.95,3.45,3.1,1.1,WHITE,LINE);text(1.5,3.75,2,.5,'Wafer map',21)
rect(4.4,3.45,3.7,1.1,WHITE,LINE);text(4.85,3.75,2.9,.5,'Yield trend',21)
text(.95,4.95,6.9,.45,'點選 die  →  側邊查看 PID / 測項',18)
text(9,2.1,3.4,.55,'一眼知道是否要看',23,bold=True)
text(9,3.3,3.4,.55,'一圖找到問題位置',23,bold=True)
text(9,4.5,3.4,.55,'一次點擊追查證據',23,bold=True)
takeaway('省下找資訊的步驟；不是在主畫面增加更多說明文字。')

slide('場景 2 的 UI：保留原始六張圖，對照同一件事','08  報表呈現','scene2/web/wafer_map_template.html；本機資料重播畫面，非現場效能證據',
      '使用原始場景二模板，保留版面與控制項。六張圖各對應一個sensor，可切預測、實測、誤差。右側MAE/散點/趨勢回答誤差大小與是否穩定。此截圖是本機重播資料，未用它宣稱線上模型驗證。兩個場景是獨立頁面，場景一採用相同視覺語言。')
img=ROOT/'reports/hc-preview/temperature-restored.png'
if img.exists():
    im=Image.open(img);im=im.crop((0,0,1440,750));im.save(OUT/'scene2_ui_crop.png')
    current.shapes.add_picture(str(OUT/'scene2_ui_crop.png'),Inches(.65),Inches(1.65),width=Inches(9.1),height=Inches(4.74))
    canvas.paste(im.resize((1092,569)),(78,198))
text(10.05,2.05,2.6,.8,'六個 sensor\n各自有誤差',20,bold=True)
text(10.05,3.55,2.6,.8,'同一顆 die\n可以交叉查看',20)
text(10.05,5,2.6,.7,'預測值 ≠ 實測值',18,ACC,True)

slide('官方系統管測試，我們補上分析與報表','09  系統分工','ONEAPI_Manual.pdf；Question 第 5–6 頁；edge_monitor.py、hc_push.py、hc_server.py',
      'ACS/Nexus透過ONEAPI提供事件與TP請求。Edge做模型運算，溫度透過官方Action回覆；異常訊息用ActionManager.set_message，TCCT輪詢取回。報表另由Edge背景HTTP POST到HC，瀏覽器查HC JSON API。AUS用於映像發布和官方部署，不是這條報表HTTP傳輸的中繼。')
for x,t,b in [(.65,'SmarTest / Nexus','執行測試\nONEAPI 事件與請求'),(4.85,'Edge 容器','場景 1 分類\n場景 2 預測'),(9.05,'HC / Browser','接收 JSON、保存報告\n兩個獨立任務頁面')]:card(x,1.9,3.65,t,b)
arrow(4.38,2.55);arrow(8.6,2.55)
text(.85,4.25,7.8,.5,'←  官方 Action：溫度回覆／異常通知',22,ACC,True)
text(.85,5.05,11.5,.65,'部署路徑：HC build → AUS registry → 官方部署機制 → Edge',21)
takeaway('模型回覆不等待報表上傳；HC 斷線不應阻塞測試 callback。')

slide('現場暴露的問題，讓可靠性成為設計的一部分','10  工程證據與限制','現場錯誤日誌；tests/test_hc_transport.py、test_task_integration.py；tasks-v3 本機修正',
      '現場曾遇到不完整JSON與BrokenPipe。舊版3秒逾時是可能因素，沒有足夠現場紀錄證明唯一根因。v3加入gzip、60秒逾時、完整接收再寫入與重試。另有晚到預測請求：讀入第一個target值前保存推論，晚到僅回傳保存值並標記late_request_frozen，排除準時預測MAE；這不能恢復已錯過的實際時機。47項本機測試通過，包括20MB壓縮payload/截斷重送。未完成Gemini現場驗收。')
card(.65,1.8,5.8,'資料傳一半斷線','壓縮＋背景重試\n完整接收後才保存')
card(6.7,1.8,6,'預測請求晚於量測','只能使用保存的前序結果\n標記延遲，不計準時成功')
text(.85,4.45,3.1,.75,'47 項',37,ACC,True);text(3.6,4.57,3.15,.7,'本機回歸測試通過',20)
text(7.05,4.45,2.65,.75,'20 MB',37,ACC,True);text(9.9,4.57,2.8,.7,'截斷與重送測試',20)
takeaway('已驗證本機失敗處理；現場根因、回覆時限與部署效果仍需驗收。')

slide('真正的價值，要用「提前多少、付出什麼代價」衡量','11  如何驗證目標','評估計畫；非已達成成效。80 die 來自題目；24 die 為示意觸發點',
      '現階段不能宣稱節省多少測試時間。示例：80顆中24顆時觸發，未測56顆=70%，只是未測die比例上限，尚未扣除誤報、停止延遲、每die測時差異等。系統目前不自動停測。需要量測：異常召回/正常誤報、首次有效通知時間、溫度準時率與p95延遲、HC資料新鮮度、UI定位任務用時。')
for x,t,b in [(.65,'場景 1','異常召回、正常誤報\n首次有效通知時間'),(4.85,'場景 2','MAE ＋ 準時回覆率\n請求到回覆的 p95 延遲'),(9.05,'操作與可靠性','找到證據所需時間\n資料新鮮度與重送成功率')]:card(x,1.8,3.65,t,b)
text(.85,4.45,11.7,.65,'示意：80 顆中第 24 顆觸發 → 尚未測 56 顆（70%）',25,ACC,True)
text(.85,5.25,11.6,.55,'這是可避免工作量的理論上限，不是已節省時間；目前沒有自動停測。',18)
takeaway('只有在誤報可接受、回覆準時且操作有效時，提前判斷才真正有價值。')

slide('設計方向有證據，量產承諾仍需下一輪驗證','12  結論','證據彙整：本簡報第 5、7、11 頁；題目第 7 頁',
      '收尾不宣稱整體完成。支持的出發點：簡單模型作為小樣本起點；合法前序訊號可改善溫度MAE；官方事件/Action與獨立背景報表是可測試架構。待證明：七類泛化、實際提前時機、現場連線穩定性、操作時間縮短、真實節省成本。')
card(.65,1.9,5.8,'已有證據支持','簡單分類器適合作為目前起點\n前序量測能改善溫度預測\n傳輸失敗情境可本機重現與測試')
card(6.7,1.9,6,'下一步才能證明','七類 wafer 的泛化能力\n現場準時回覆與穩定運行\n人員操作時間與實際節省')
text(.85,4.6,11.6,.95,'把正確資料，在正確時機，交給需要做決定的人與程式。',28,ACC,True)

slide('附錄｜數字與設計的來源','EVIDENCE & REPRODUCIBILITY','專案內部簡報；原始競賽資料不隨此簡報附送',
      '題目引用以20260919版本為準。舊README部分敘述（尚未替換即時服務）已過時，本簡報模型現況以目前程式為準。評估數字以comparison.json/evaluation.json為準。數據證据json附在presentations/evidence.json。不要將in-sample24/24、未校準confidence或本機47tests當作量產准确率。')
sources=[('問題與限制','doc/Question_20260919.pdf：第 1–7 頁'),('官方介面','doc/ONEAPI_Manual.pdf；SmarTest Main.flow / Predict.java'),('場景 1 選型','artifacts/wafer_classifier/info.json；reports/wafer_classifier/comparison.json'),('場景 2 效能','artifacts/scene2/evaluation.json；presentations/evidence.json'),('因果與傳輸測試','tests/test_task_integration.py；tests/test_hc_transport.py'),('部署與版本','deploy/HC_EDGE_TASKS.md；tasks-v3 為本機修正版，待現場驗收')]
for i,(a,b) in enumerate(sources):
    y=1.8+i*.68;text(.75,y,2.1,.45,a,18,ACC,True);text(3,y,9.7,.55,b,16)

pages.append(canvas)
OUT.mkdir(exist_ok=True)
prs.save(OUT/'ACS_RTDI_design_story.pptx')
pages[0].save(OUT/'ACS_RTDI_design_story.pdf',save_all=True,append_images=pages[1:],resolution=150)
thumb=Image.new('RGB',(1200,math.ceil(len(pages)/3)*250),'#e4e7e9')
for i,p in enumerate(pages):
    p.save(OUT/('slide_%02d.png'%(i+1)))
    thumb.paste(p.resize((384,216)),(8+(i%3)*400,8+(i//3)*250))
thumb.save(OUT/'overview.png')
(OUT/'layout_check.json').write_text(json.dumps({'slides':len(pages),'text_overflow_checks':overflows},ensure_ascii=False,indent=2),encoding='utf-8')
print('Created',len(pages),'slides. Text overflow checks:',overflows)
